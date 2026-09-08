"""Offline QA suite - 34 assertions. Works with or without a Gemini key.

Run:  ./.venv/Scripts/python.exe tests/qa_offline.py
Exit code 0 = all passed.
"""
import io
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import numpy as np
from PIL import Image
from fastapi.testclient import TestClient

import app.main
from app.services.normalize import normalize_email, normalize_phone
from app.services.excel_store import _safe_cell

c = TestClient(app.main.app)
PASS, FAIL = [], []


def check(name, cond, detail=""):
    (PASS if cond else FAIL).append(name)
    print(f"  [{'PASS' if cond else 'FAIL'}] {name}" + (f"  -- {detail}" if detail else ""))


def jpeg(w, h, noisy=True):
    src = (np.random.rand(h, w, 3) * 255).astype("uint8") if noisy else np.full((h, w, 3), 128, "uint8")
    b = io.BytesIO()
    Image.fromarray(src).save(b, format="JPEG", quality=92)
    return b.getvalue()


print("\n== 1. normalize_phone ==")
for raw, expect in {
    "+91 98765 43210": "+91 98765 43210",
    "+91 80 4123 9876  |  +91 99860 45521": "+91 80412 39876, +91 99860 45521",
    "9876543210/09123456789": "+91 98765 43210, +91 91234 56789",
    "": "",
    "no digits here": "",
}.items():
    got, _ = normalize_phone(raw)
    check(f"phone {raw!r}", got == expect, f"got={got!r}")
for raw in ("+1 (415) 555-2671", "98765 43210 ext 22"):
    _, clean = normalize_phone(raw)
    check(f"phone {raw!r} flagged not-clean", clean is False)

print("\n== 2. normalize_email ==")
check("dedupe+lowercase", normalize_email(" A@B.com , a@b.com") == "a@b.com")
check("empty", normalize_email("") == "")
check("multi joined", normalize_email("x@y.com; z@y.com") == "x@y.com, z@y.com")

print("\n== 3. _safe_cell (Excel hardening) ==")
check("formula = neutralised", _safe_cell("=1+1").startswith("'="))
check("formula + neutralised", _safe_cell("+91 98765 43210").startswith("'+"))
check("formula @ neutralised", _safe_cell("@SUM(A1)").startswith("'@"))
check("plain text untouched", _safe_cell("Rahul Sharma") == "Rahul Sharma")
check("control chars stripped", _safe_cell("a\x00b\x07c") == "abc")
check("over-long truncated", len(_safe_cell("x" * 40000)) <= 32767)
check("None -> blank", _safe_cell(None) == "")

print("\n== 4. /api/extract edge inputs ==")
check("empty file -> 400", c.post("/api/extract", files={"image": ("x.jpg", b"", "image/jpeg")}).status_code == 400)
check("non-image -> 400", c.post("/api/extract", files={"image": ("x.pdf", b"%PDF-1.4 nope", "application/pdf")}).status_code == 400)
check("corrupt jpeg -> 400", c.post("/api/extract", files={"image": ("x.jpg", b"\xff\xd8\xff\xff junk", "image/jpeg")}).status_code == 400)
check("oversize -> 413", c.post("/api/extract", files={"image": ("x.jpg", b"\xff\xd8" + b"0" * (13 * 1024 * 1024), "image/jpeg")}).status_code == 413)
check("missing 'image' field -> 422", c.post("/api/extract").status_code == 422)
rj = c.post("/api/extract", files={"image": ("x.jpg", jpeg(800, 500, noisy=False), "image/jpeg")}).json()
check("flat gray -> blurry:true", rj.get("blurry") is True)
check("tiny 24x16 -> no 500", c.post("/api/extract", files={"image": ("x.jpg", jpeg(24, 16), "image/jpeg")}).status_code == 200)
b = io.BytesIO()
Image.fromarray((np.random.rand(400, 700, 4) * 255).astype("uint8"), "RGBA").save(b, "PNG")
check("RGBA PNG accepted", c.post("/api/extract", files={"image": ("x.png", b.getvalue(), "image/png")}).status_code == 200)

print("\n== 5. /api/leads validation + hardening ==")
base = dict(name="A", company="B", title="C", email="a@b.com", phone="+91 98765 43210", notes="n", assignedTo="NITISH")
check("email without TLD -> 422", c.post("/api/leads", json={**base, "email": "a@b"}).status_code == 422)
check("phone no digits -> 422", c.post("/api/leads", json={**base, "phone": "abcd"}).status_code == 422)
check("wrong-case assignee -> 422", c.post("/api/leads", json={**base, "assignedTo": "nitish"}).status_code == 422)
check("blank notes -> 422", c.post("/api/leads", json={**base, "notes": ""}).status_code == 422)
check("null field -> 422 not 500", c.post("/api/leads", json={**base, "name": None}).status_code == 422)
r = c.post("/api/leads", json={**base, "name": "=cmd|' /c calc'!A1", "notes": "x\x01y"})
check("injection + control chars save OK", r.status_code == 200, r.text[:120])
if r.status_code == 200:
    from openpyxl import load_workbook
    ws = load_workbook(app.main.settings.excel_path).active
    last = [x.value for x in ws[ws.max_row]]
    check("  stored name quote-guarded", str(last[0]).startswith("'="), f"stored={last[0]!r}")
    check("  control char stripped", "\x01" not in str(last[5]))
check("50k-char notes save OK", c.post("/api/leads", json={**base, "notes": "z" * 50000}).status_code == 200)

# tidy the test workbook
try:
    os.remove(app.main.settings.excel_path)
    os.rmdir(os.path.dirname(app.main.settings.excel_path))
except OSError:
    pass

print(f"\n==== {len(PASS)} passed, {len(FAIL)} failed ====")
sys.exit(1 if FAIL else 0)
